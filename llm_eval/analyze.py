"""Unblind the scored workbook and produce summary.json, report.md and chart.svg."""
from collections import Counter
from html import escape
import json
from pathlib import Path
import re
import statistics

from openpyxl import load_workbook

from .rubric import CATEGORIES, ERROR_TYPES, parse_error

MARK_START = "<!-- RESULTS:START -->"
MARK_END = "<!-- RESULTS:END -->"
COLORS = ["#2F6FB0", "#E07B39", "#3A9E6F", "#8C5BB5", "#C0504D"]


def read_scores(path: Path) -> dict[str, dict]:
    ws = load_workbook(path, read_only=True, data_only=True)["打分"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    col = {name: header.index(name) for name in ("编号", "分数(1-5)", "主要错误类型", "次要错误类型", "评分理由")}
    scores = {}
    for values in rows:
        item = values[col["编号"]]
        if not item:
            continue
        raw = values[col["分数(1-5)"]]
        score = None
        if raw not in (None, ""):
            score = float(raw)
            if score != int(score) or not 1 <= score <= 5:
                raise ValueError(f"{item}: score must be an integer from 1 to 5, got {raw}")
            score = int(score)
        errors = [e for e in (parse_error(values[col["主要错误类型"]]), parse_error(values[col["次要错误类型"]])) if e]
        scores[str(item)] = {"score": score, "errors": list(dict.fromkeys(errors)),
                             "rationale": str(values[col["评分理由"]] or "").strip()}
    return scores


def _mean(values):
    return round(statistics.fmean(values), 2) if values else None


def analyze(questions: list[dict], key: dict, scores: dict, rescore: dict | None = None) -> dict:
    items = key["items"]
    unknown = set(scores) - set(items)
    if unknown:
        raise ValueError(f"Workbook has items not in key.json: {sorted(unknown)[:5]}")
    category = {q["id"]: q["category"] for q in questions}
    models = key["models"]
    scored = {i: s for i, s in scores.items() if s["score"] is not None}
    per_model = {}
    for m in models:
        mine = [(i, s) for i, s in scored.items() if items[i] == m]
        values = [s["score"] for _, s in mine]
        per_model[m] = {
            "scored": len(values),
            "mean": _mean(values),
            "by_category": {c: _mean([s["score"] for i, s in mine if category[i.rsplit("-", 1)[0]] == c]) for c in CATEGORIES},
            "distribution": {str(k): values.count(k) for k in range(5, 0, -1)},
            "full_marks_rate": round(values.count(5) / len(values), 4) if values else None,
            "errors": dict(Counter(e for _, s in mine for e in s["errors"]).most_common()),
        }
    per_question = []
    for q in questions:
        row = {m: next((s["score"] for i, s in scored.items() if i.rsplit("-", 1)[0] == q["id"] and items[i] == m), None)
               for m in models}
        present = [v for v in row.values() if v is not None]
        per_question.append({"id": q["id"], "category": q["category"], "scores": row,
                             "spread": max(present) - min(present) if len(present) > 1 else None})
    summary = {"models": models, "items": len(items), "scored_items": len(scored),
               "coverage": round(len(scored) / len(items), 4) if items else 0,
               "per_model": per_model,
               "error_totals": dict(Counter(e for s in scored.values() for e in s["errors"]).most_common()),
               "most_divergent": sorted((q for q in per_question if q["spread"] is not None),
                                        key=lambda q: (-q["spread"], q["id"]))[:5],
               "per_question": per_question}
    if rescore:
        pairs = [(scored[i]["score"], r["score"]) for i, r in rescore.items() if r["score"] is not None and i in scored]
        if pairs:
            summary["consistency"] = {"items": len(pairs),
                                      "exact_agreement": round(sum(a == b for a, b in pairs) / len(pairs), 4),
                                      "within_one": round(sum(abs(a - b) <= 1 for a, b in pairs) / len(pairs), 4),
                                      "mean_abs_diff": round(statistics.fmean(abs(a - b) for a, b in pairs), 2)}
    return summary


def _fmt(v, pct=False):
    if v is None:
        return "—"
    return f"{v:.0%}" if pct else f"{v:.2f}"


def render_markdown(s: dict, chart: str | None = "chart.svg", note: str = "") -> str:
    lines = []
    if note:
        lines += [note, ""]
    lines += [f"已打分 {s['scored_items']} / {s['items']} 条回答（{s['coverage']:.0%}）。分数 1–5，越高越好。", ""]
    if chart:
        lines += [f"![各模型分类别平均分]({chart})", ""]
    head = "| 模型 | 总平均分 | " + " | ".join(CATEGORIES) + " | 满分率 | 主要错误类型 |"
    lines += [head, "| --- | ---: | " + " | ".join("---:" for _ in CATEGORIES) + " | ---: | --- |"]
    for m in s["models"]:
        p = s["per_model"][m]
        errs = "、".join(f"{ERROR_TYPES[e][0]} {n}" for e, n in list(p["errors"].items())[:3]) or "—"
        lines.append(f"| {m} | {_fmt(p['mean'])} | " + " | ".join(_fmt(p["by_category"][c]) for c in CATEGORIES)
                     + f" | {_fmt(p['full_marks_rate'], True)} | {errs} |")
    if s["error_totals"]:
        lines += ["", "**错误类型分布**（全部模型合计）", "", "| 错误类型 | 次数 |", "| --- | ---: |"]
        lines += [f"| {ERROR_TYPES[e][0]} | {n} |" for e, n in s["error_totals"].items()]
    if s["most_divergent"]:
        lines += ["", "**模型之间分歧最大的题**", "", "| 题号 | 类别 | " + " | ".join(s["models"]) + " | 分差 |",
                  "| --- | --- | " + " | ".join("---:" for _ in s["models"]) + " | ---: |"]
        for q in s["most_divergent"]:
            lines.append(f"| {q['id']} | {q['category']} | " + " | ".join(str(q["scores"][m] or "—") for m in s["models"])
                         + f" | {q['spread']} |")
    if "consistency" in s:
        c = s["consistency"]
        lines += ["", f"**自我复评一致性**：隔天随机抽 {c['items']} 条重新打分，完全一致 {c['exact_agreement']:.0%}，"
                      f"相差不超过 1 分 {c['within_one']:.0%}，平均分差 {c['mean_abs_diff']}。"]
    return "\n".join(lines) + "\n"


def render_svg(s: dict, title: str) -> str:
    groups = CATEGORIES + ["总平均"]
    models = s["models"]
    width, height, left, top, bottom = 760, 340, 48, 54, 300
    group_w = (width - left - 20) / len(groups)
    bar_w = min(26, (group_w - 24) / max(len(models), 1))
    y = lambda v: bottom - (bottom - top) * (v / 5)
    parts = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height + 40}" viewBox="0 0 {width} {height + 40}" '
             'font-family="-apple-system,BlinkMacSystemFont,\'PingFang SC\',\'Microsoft YaHei\',sans-serif">',
             f'<rect width="100%" height="100%" fill="#ffffff"/>',
             f'<text x="{left}" y="28" font-size="16" font-weight="600" fill="#1F2937">{escape(title)}</text>']
    for v in range(0, 6):
        parts.append(f'<line x1="{left}" x2="{width - 20}" y1="{y(v):.1f}" y2="{y(v):.1f}" stroke="#E5E7EB"/>')
        parts.append(f'<text x="{left - 8}" y="{y(v) + 4:.1f}" font-size="11" text-anchor="end" fill="#6B7280">{v}</text>')
    for gi, g in enumerate(groups):
        gx = left + gi * group_w + (group_w - bar_w * len(models)) / 2
        for mi, m in enumerate(models):
            p = s["per_model"][m]
            v = p["mean"] if g == "总平均" else p["by_category"][g]
            if v is None:
                continue
            x = gx + mi * bar_w
            parts.append(f'<rect x="{x:.1f}" y="{y(v):.1f}" width="{bar_w - 3:.1f}" height="{bottom - y(v):.1f}" '
                         f'rx="2" fill="{COLORS[mi % len(COLORS)]}"/>')
            parts.append(f'<text x="{x + (bar_w - 3) / 2:.1f}" y="{y(v) - 5:.1f}" font-size="10" text-anchor="middle" '
                         f'fill="#374151">{v:.1f}</text>')
        parts.append(f'<text x="{left + gi * group_w + group_w / 2:.1f}" y="{bottom + 20}" font-size="12" '
                     f'text-anchor="middle" fill="#374151">{g}</text>')
    lx = left
    for mi, m in enumerate(models):
        parts.append(f'<rect x="{lx}" y="{height + 8}" width="12" height="12" rx="2" fill="{COLORS[mi % len(COLORS)]}"/>')
        parts.append(f'<text x="{lx + 18}" y="{height + 18}" font-size="12" fill="#374151">{escape(m)}</text>')
        lx += 30 + 13 * len(m)
    parts.append("</svg>")
    return "\n".join(parts) + "\n"


def update_readme(readme: Path, block: str) -> None:
    text = readme.read_text(encoding="utf-8")
    pattern = re.compile(re.escape(MARK_START) + r".*?" + re.escape(MARK_END), re.DOTALL)
    if not pattern.search(text):
        raise ValueError(f"{readme} has no {MARK_START} ... {MARK_END} markers")
    readme.write_text(pattern.sub(lambda _: f"{MARK_START}\n{block.rstrip()}\n{MARK_END}", text), encoding="utf-8")


def run_analysis(questions: list[dict], run_dir: Path, title: str, note: str = "") -> dict:
    key = json.loads((run_dir / "key.json").read_text(encoding="utf-8"))
    rescore_path = run_dir / "rescore.xlsx"
    rescore = read_scores(rescore_path) if rescore_path.exists() else None
    summary = analyze(questions, key, read_scores(run_dir / "scoring.xlsx"), rescore)
    (run_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (run_dir / "chart.svg").write_text(render_svg(summary, title), encoding="utf-8")
    (run_dir / "report.md").write_text(f"# {title}\n\n" + render_markdown(summary, "chart.svg", note), encoding="utf-8")
    return summary
