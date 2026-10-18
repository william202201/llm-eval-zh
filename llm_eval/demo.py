"""End-to-end demo with clearly synthetic answers and scores, to show the workflow and report format.

Nothing here is a real model evaluation: the "models" are placeholders and the scores are random.
"""
import json
from pathlib import Path
import random

from openpyxl import load_workbook

from .analyze import run_analysis
from .rubric import ERROR_LABELS
from .sheet import build_rescore, build_sheet

DEMO_MODELS = {"示例模型A": 0.6, "示例模型B": 0.0, "示例模型C": -0.5}
CATEGORY_ERRORS = {"推理": "reasoning", "写作": "instruction", "事实": "hallucination", "安全": "over_refusal"}
DEMO_NOTE = "> 演示数据：回答是占位文本，分数由固定随机种子生成，只用于展示流程和报告格式，不代表任何真实模型的表现。"


def _label(code):
    return next(label for label in ERROR_LABELS if label.startswith(code + "｜"))


def _fill_scores(path: Path, key: dict, questions: list[dict], rng: random.Random, jitter: float = 0.0) -> None:
    category = {q["id"]: q["category"] for q in questions}
    wb = load_workbook(path)
    ws = wb["打分"]
    for row in ws.iter_rows(min_row=2):
        item = row[0].value
        if not item:
            continue
        bias = DEMO_MODELS[key["items"][item]]
        score = max(1, min(5, round(rng.gauss(3.9 + bias, 0.9) + jitter * rng.choice([-1, 0, 0, 1]))))
        row[6].value = score
        if score <= 3:
            row[7].value = _label(CATEGORY_ERRORS[category[item.rsplit("-", 1)[0]]])
            if score == 1 and rng.random() < 0.5:
                row[8].value = _label("verbose")
            row[9].value = "（演示）扣分理由示例"
        else:
            row[9].value = "（演示）满足要求"
    wb.save(path)


def run_demo(questions: list[dict], out_dir: Path, seed: int = 2026) -> dict:
    rng = random.Random(seed)
    # Placeholder text is identical across models so the demo sheet stays blind, like a real run.
    answers = {m: {q["id"]: f"（演示占位）第 {q['id']} 题的回答。真实评测时这里是模型的原始输出。" for q in questions}
               for m in DEMO_MODELS}
    build_sheet(questions, answers, out_dir, seed)
    key = json.loads((out_dir / "key.json").read_text(encoding="utf-8"))
    _fill_scores(out_dir / "scoring.xlsx", key, questions, rng)
    # Simulated re-score of 10 items: copy the original scores with occasional ±1 drift.
    build_rescore(out_dir, 10)
    original = {r[0].value: r[6].value for r in load_workbook(out_dir / "scoring.xlsx")["打分"].iter_rows(min_row=2)}
    wb = load_workbook(out_dir / "rescore.xlsx")
    for row in wb["打分"].iter_rows(min_row=2):
        if row[0].value:
            row[6].value = max(1, min(5, original[row[0].value] + rng.choice([0, 0, 0, 1, -1])))
    wb.save(out_dir / "rescore.xlsx")
    return run_analysis(questions, out_dir, "中文大模型评测报告（演示数据）", DEMO_NOTE)
