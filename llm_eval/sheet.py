"""Build the blind scoring workbook. The model behind each answer lives only in key.json."""
import hashlib
import json
import math
from pathlib import Path
import random

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .rubric import CATEGORY_FOCUS, ERROR_LABELS, ERROR_TYPES, SCORE_SCALE

HEADERS = ["编号", "题号", "类别", "问题", "参考要点", "回答", "分数(1-5)", "主要错误类型", "次要错误类型", "评分理由"]
WIDTHS = [9, 6, 6, 34, 30, 60, 10, 20, 20, 36]
WRAP_COLUMNS = {4, 5, 6, 10}
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
GROUP_FILL = PatternFill("solid", fgColor="F3F6FA")


def _height(texts_and_widths) -> float:
    lines = 1
    for text, width in texts_and_widths:
        text = str(text or "")
        # CJK characters are roughly twice as wide as the column-width unit.
        est = sum(math.ceil(max(1, len(part) * 2) / max(width, 1)) for part in text.split("\n"))
        lines = max(lines, est)
    return min(409, 15 * lines + 4)


def _style(ws):
    for col, (title, width) in enumerate(zip(HEADERS, WIDTHS), start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "D2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def _validations(ws, last_row):
    score = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                           showErrorMessage=True, errorTitle="分数无效", error="请输入 1 到 5 的整数")
    errors = DataValidation(type="list", formula1="=选项!$A$1:$A$%d" % len(ERROR_LABELS), allow_blank=True)
    ws.add_data_validation(score)
    ws.add_data_validation(errors)
    score.add(f"G2:G{last_row}")
    errors.add(f"H2:I{last_row}")


def _options_sheet(wb):
    ws = wb.create_sheet("选项")
    for i, label in enumerate(ERROR_LABELS, start=1):
        ws.cell(row=i, column=1, value=label)
    ws.column_dimensions["A"].width = 28
    ws.sheet_state = "hidden"


def _rubric_sheet(wb):
    ws = wb.create_sheet("评分标准")
    rows = [("分数", "含义")] + [(str(k), v) for k, v in sorted(SCORE_SCALE.items(), reverse=True)]
    rows += [("", ""), ("类别", "重点看什么")] + list(CATEGORY_FOCUS.items())
    rows += [("", ""), ("错误类型", "判定说明")] + [(f"{c}｜{n}", d) for c, (n, d) in ERROR_TYPES.items()]
    rows += [("", ""), ("提示", "每道题的几个回答顺序是随机的，编号不代表模型。打完全部分数前不要打开 key.json。")]
    for r, (a, b) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=a)
        ws.cell(row=r, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
        if b in ("含义", "重点看什么", "判定说明"):
            ws.cell(row=r, column=1).font = ws.cell(row=r, column=2).font = Font(bold=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 80


def build_sheet(questions: list[dict], answers: dict[str, dict[str, str]], out_dir: Path, seed: int = 2026) -> dict:
    """answers: {model_name: {question_id: answer}}. Writes scoring.xlsx and key.json into out_dir."""
    if len(answers) < 2:
        raise ValueError("Need answers from at least 2 models for a comparison")
    missing = {m: [q["id"] for q in questions if q["id"] not in a] for m, a in answers.items()}
    missing = {m: ids for m, ids in missing.items() if ids}
    if missing:
        raise ValueError(f"Missing answers: {missing}")
    out_dir.mkdir(parents=True, exist_ok=True)
    rng = random.Random(seed)
    wb = Workbook()
    ws = wb.active
    ws.title = "打分"
    _style(ws)
    key, row = {}, 2
    models = sorted(answers)
    for qi, q in enumerate(questions):
        order = models[:]
        rng.shuffle(order)
        reference = q["reference"] + "\n检查：" + "；".join(q["checks"])
        for n, model in enumerate(order, start=1):
            item = f"{q['id']}-{n}"
            key[item] = model
            values = [item, q["id"], q["category"], q["prompt"], reference, answers[model][q["id"]], None, None, None, None]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.alignment = Alignment(wrap_text=col in WRAP_COLUMNS, vertical="top")
                if qi % 2:
                    cell.fill = GROUP_FILL
            ws.row_dimensions[row].height = _height([(q["prompt"], WIDTHS[3]), (reference, WIDTHS[4]),
                                                     (answers[model][q["id"]], WIDTHS[5])])
            row += 1
    _validations(ws, row - 1)
    _rubric_sheet(wb)
    _options_sheet(wb)
    wb.save(out_dir / "scoring.xlsx")
    fingerprint = hashlib.sha256(json.dumps([q["id"] for q in questions]).encode()).hexdigest()[:12]
    meta = {"seed": seed, "models": models, "questions": len(questions), "question_ids_sha256": fingerprint, "items": key}
    (out_dir / "key.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return {"items": len(key), "models": models, "workbook": str(out_dir / "scoring.xlsx")}


def build_rescore(run_dir: Path, n: int = 10, seed: int = 7) -> dict:
    """Copy a random sample of items into rescore.xlsx with blank scores, for a self-consistency check days later."""
    source = load_workbook(run_dir / "scoring.xlsx")["打分"]
    rows = [[c.value for c in r] for r in source.iter_rows(min_row=2) if r[0].value]
    if n < 1 or n > len(rows):
        raise ValueError(f"n must be between 1 and {len(rows)}")
    sample = random.Random(seed).sample(rows, n)
    wb = Workbook()
    ws = wb.active
    ws.title = "打分"
    _style(ws)
    for r, values in enumerate(sample, start=2):
        for col, value in enumerate(values[:6] + [None] * 4, start=1):
            ws.cell(row=r, column=col, value=value).alignment = Alignment(wrap_text=col in WRAP_COLUMNS, vertical="top")
        ws.row_dimensions[r].height = _height([(values[3], WIDTHS[3]), (values[4], WIDTHS[4]), (values[5], WIDTHS[5])])
    _validations(ws, n + 1)
    _rubric_sheet(wb)
    _options_sheet(wb)
    wb.save(run_dir / "rescore.xlsx")
    return {"items": n, "workbook": str(run_dir / "rescore.xlsx")}
