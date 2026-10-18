import csv
from http.server import BaseHTTPRequestHandler, HTTPServer
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import threading
import unittest
from unittest import mock

from openpyxl import load_workbook

from llm_eval.analyze import analyze, read_scores, render_markdown, update_readme
from llm_eval.cli import DEFAULT_QUESTIONS, main
from llm_eval.collect import collect, load_models
from llm_eval.data import load_answers, load_questions, write_template
from llm_eval.rubric import ERROR_LABELS
from llm_eval.sheet import build_rescore, build_sheet


class EvalTest(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.questions = load_questions(DEFAULT_QUESTIONS)[:4]
        self.answers = {m: {q["id"]: f"{m} answer {q['id']}" for q in self.questions} for m in ("alpha", "beta", "gamma")}

    def score(self, path, scores):
        wb = load_workbook(path)
        for row in wb["打分"].iter_rows(min_row=2):
            if row[0].value in scores:
                value = scores[row[0].value]
                row[6].value, row[7].value = value if isinstance(value, tuple) else (value, None)
        wb.save(path)

    def test_bundled_question_set_is_valid(self):
        questions = load_questions(DEFAULT_QUESTIONS)
        self.assertEqual(len(questions), 30)
        self.assertEqual({q["category"] for q in questions}, {"推理", "写作", "事实", "安全"})

    def test_question_validation(self):
        path = self.root / "q.jsonl"
        row = {"id": "X1", "category": "推理", "prompt": "p", "reference": "r", "checks": ["c"]}
        path.write_text(json.dumps(row) + "\n" + json.dumps(row), encoding="utf-8")
        with self.assertRaises(ValueError):
            load_questions(path)
        path.write_text(json.dumps({**row, "category": "数学"}), encoding="utf-8")
        with self.assertRaises(ValueError):
            load_questions(path)

    def test_sheet_is_blind_and_key_maps_every_item(self):
        out = self.root / "run"
        result = build_sheet(self.questions, self.answers, out, seed=1)
        self.assertEqual(result["items"], 12)
        key = json.loads((out / "key.json").read_text(encoding="utf-8"))
        ws = load_workbook(out / "scoring.xlsx")["打分"]
        cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        for model in self.answers:
            self.assertFalse(any(cell == model for cell in cells), "model name must not appear in the workbook")
        for row in ws.iter_rows(min_row=2):
            item, answer = row[0].value, row[5].value
            self.assertTrue(answer.startswith(key["items"][item]))
        orders = [[key["items"][f"{q['id']}-{n}"] for n in (1, 2, 3)] for q in self.questions]
        self.assertGreater(len({tuple(o) for o in orders}), 1, "answer order should be shuffled per question")

    def test_missing_answers_and_single_model_rejected(self):
        broken = {**self.answers, "beta": {}}
        with self.assertRaises(ValueError):
            build_sheet(self.questions, broken, self.root / "a")
        with self.assertRaises(ValueError):
            build_sheet(self.questions, {"alpha": self.answers["alpha"]}, self.root / "b")

    def test_analysis_unblinds_and_aggregates(self):
        out = self.root / "run"
        build_sheet(self.questions, self.answers, out, seed=1)
        key = json.loads((out / "key.json").read_text(encoding="utf-8"))
        target = {"alpha": 5, "beta": 3, "gamma": 1}
        reasoning = next(label for label in ERROR_LABELS if label.startswith("reasoning"))
        scores = {item: (target[m], reasoning if target[m] < 3 else None) for item, m in key["items"].items()}
        scores["R01-1"] = None  # leave one unscored
        self.score(out / "scoring.xlsx", {k: v for k, v in scores.items() if v is not None})
        summary = analyze(self.questions, key, read_scores(out / "scoring.xlsx"))
        self.assertEqual(summary["scored_items"], 11)
        for model, value in target.items():
            self.assertEqual(summary["per_model"][model]["mean"], value)
        self.assertEqual(summary["per_model"]["gamma"]["errors"], {"reasoning": summary["per_model"]["gamma"]["scored"]})
        self.assertEqual(summary["most_divergent"][0]["spread"], 4)
        self.assertIn("| alpha | 5.00 |", render_markdown(summary))

    def test_invalid_score_is_rejected(self):
        out = self.root / "run"
        build_sheet(self.questions, self.answers, out, seed=1)
        self.score(out / "scoring.xlsx", {"R01-1": 7})
        with self.assertRaises(ValueError):
            read_scores(out / "scoring.xlsx")

    def test_rescore_consistency(self):
        out = self.root / "run"
        build_sheet(self.questions, self.answers, out, seed=1)
        key = json.loads((out / "key.json").read_text(encoding="utf-8"))
        self.score(out / "scoring.xlsx", {item: 4 for item in key["items"]})
        build_rescore(out, 4)
        items = [r[0].value for r in load_workbook(out / "rescore.xlsx")["打分"].iter_rows(min_row=2)]
        self.score(out / "rescore.xlsx", {items[0]: 4, items[1]: 4, items[2]: 5, items[3]: 2})
        summary = analyze(self.questions, key, read_scores(out / "scoring.xlsx"), read_scores(out / "rescore.xlsx"))
        self.assertEqual(summary["consistency"], {"items": 4, "exact_agreement": 0.5, "within_one": 0.75, "mean_abs_diff": 0.75})

    def test_readme_markers(self):
        readme = self.root / "README.md"
        readme.write_text("intro\n<!-- RESULTS:START -->\nold\n<!-- RESULTS:END -->\nend\n", encoding="utf-8")
        update_readme(readme, "new table")
        self.assertEqual(readme.read_text(encoding="utf-8"), "intro\n<!-- RESULTS:START -->\nnew table\n<!-- RESULTS:END -->\nend\n")
        readme.write_text("no markers", encoding="utf-8")
        with self.assertRaises(ValueError):
            update_readme(readme, "x")

    def test_csv_template_round_trip(self):
        path = self.root / "answers" / "web.csv"
        write_template(self.questions, path)
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        rows[1][3] = "粘贴的回答"
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            csv.writer(handle).writerows(rows)
        self.assertEqual(load_answers(path), {self.questions[0]["id"]: "粘贴的回答"})
        self.assertEqual(main(["template", "web", "--out", str(path.parent)]), 2, "existing template must not be overwritten")

    def test_collect_resumes_and_records_failures(self):
        calls = []

        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                calls.append((self.headers["Authorization"], body["model"]))
                payload = json.dumps({"choices": [{"message": {"content": f"echo {len(calls)}"}}],
                                      "usage": {"total_tokens": 3}}).encode()
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(payload)

            def log_message(self, *args):
                pass

        server = HTTPServer(("127.0.0.1", 0), Handler)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        config = self.root / "models.json"
        config.write_text(json.dumps([{"name": "fake", "base_url": f"http://127.0.0.1:{server.server_port}/v1",
                                       "model": "fake-1", "api_key_env": "FAKE_KEY"}]), encoding="utf-8")
        model = load_models(config)[0]
        out = self.root / "answers" / "fake.jsonl"
        with mock.patch.dict(os.environ, {}, clear=True):
            stats = collect(self.questions, model, out)
        self.assertEqual((stats["ok"], stats["failed"]), (0, 4))
        with mock.patch.dict(os.environ, {"FAKE_KEY": "secret"}):
            stats = collect(self.questions, model, out)
            self.assertEqual((stats["ok"], stats["failed"], stats["skipped"]), (4, 0, 0))
            stats = collect(self.questions, model, out)
            self.assertEqual((stats["ok"], stats["skipped"]), (0, 4))
        self.assertEqual(calls[0], ("Bearer secret", "fake-1"))
        self.assertEqual(len(load_answers(out)), 4)

    def test_cli_demo_end_to_end(self):
        out, docs = self.root / "demo", self.root / "docs"
        self.assertEqual(main(["demo", "--out", str(out), "--docs", str(docs)]), 0)
        for name in ("scoring.xlsx", "rescore.xlsx", "key.json", "summary.json", "report.md", "chart.svg"):
            self.assertTrue((out / name).exists(), name)
        self.assertIn("演示数据", (docs / "demo-report.md").read_text(encoding="utf-8"))
        self.assertIn("demo-chart.svg", (docs / "demo-report.md").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
